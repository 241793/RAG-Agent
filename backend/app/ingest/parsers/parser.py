"""文档解析器：统一输出 (text, pages) 结构。

Pages 结构：[{"page": 1, "text": "..."}]，供页码溯源。
支持：pdf / docx / xlsx / pptx / md / html / csv / txt。
"""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class ParsedDoc:
    text: str
    pages: list[dict] = field(default_factory=list)  # [{"page":int, "text":str}]
    page_count: int = 0
    char_count: int = 0


class ParseError(Exception):
    pass


def parse_file(path: Path, ext: str | None = None) -> ParsedDoc:
    ext = (ext or path.suffix).lower().lstrip(".")
    try:
        if ext == "pdf":
            return _parse_pdf(path)
        if ext == "docx":
            return _parse_docx(path)
        if ext in ("xlsx", "xls"):
            return _parse_xlsx(path)
        if ext == "pptx":
            return _parse_pptx(path)
        if ext in ("md", "markdown"):
            return _parse_markdown(path)
        if ext in ("html", "htm"):
            return _parse_html(path)
        if ext == "csv":
            return _parse_csv(path)
        if ext in ("txt", "text", "log"):
            return _parse_text(path)
    except ParseError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ParseError(f"解析 {ext} 失败: {e}") from e
    return _parse_text(path)


def _finalize(pages: list[dict]) -> ParsedDoc:
    text = "\n\n".join(p["text"] for p in pages if p.get("text"))
    return ParsedDoc(text=text, pages=pages, page_count=len(pages), char_count=len(text))


def _parse_pdf(path: Path) -> ParsedDoc:
    import fitz  # PyMuPDF

    pages: list[dict] = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc):
            pages.append({"page": i + 1, "text": page.get_text("text").strip()})
    return _finalize(pages)


def _parse_docx(path: Path) -> ParsedDoc:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            style = (para.style.name or "").lower()
            if "heading" in style or "标题" in style:
                parts.append(f"\n## {para.text.strip()}\n")
            else:
                parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return _finalize([{"page": 1, "text": "\n".join(parts)}])


def _parse_xlsx(path: Path) -> ParsedDoc:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    pages: list[dict] = []
    for idx, ws in enumerate(wb.worksheets):
        lines: list[str] = [f"# 工作表: {ws.title}"]
        header: list[str] = []
        for row_i, row in enumerate(ws.iter_rows(values_only=True)):
            # 保留列位置（空单元格补空串），避免列错位——表格问答依赖列对齐
            cells = ["" if c is None else str(c) for c in row]
            # 去掉右侧整列为空的占位
            while cells and cells[-1] == "":
                cells.pop()
            if not any(cells):
                continue
            if row_i == 0:
                header = cells
            lines.append(" | ".join(cells))
        if header:
            # 顶部插入表头标注，便于 AI 识别列名
            lines.insert(1, f"列名: {' | '.join(header)}")
        pages.append({"page": idx + 1, "text": "\n".join(lines)})
    wb.close()
    return _finalize(pages)


def _parse_pptx(path: Path) -> ParsedDoc:
    from pptx import Presentation

    prs = Presentation(str(path))
    pages: list[dict] = []
    for i, slide in enumerate(prs.slides):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        texts.append(t)
        pages.append({"page": i + 1, "text": "\n".join(texts)})
    return _finalize(pages)


def _parse_markdown(path: Path) -> ParsedDoc:
    return _finalize([{"page": 1, "text": path.read_text(encoding="utf-8", errors="ignore")}])


def _parse_html(path: Path) -> ParsedDoc:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "lxml")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    return _finalize([{"page": 1, "text": soup.get_text("\n", strip=True)}])


def _parse_csv(path: Path) -> ParsedDoc:
    content = _decode(path.read_bytes())
    reader = csv.reader(io.StringIO(content))
    lines = [" | ".join(row) for row in reader if any(c.strip() for c in row)]
    return _finalize([{"page": 1, "text": "\n".join(lines)}])


def _parse_text(path: Path) -> ParsedDoc:
    return _finalize([{"page": 1, "text": _decode(path.read_bytes())}])


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gbk", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")
