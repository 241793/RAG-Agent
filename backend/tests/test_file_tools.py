"""文件处理工具测试：生成/转换产物 + Artifact 落库 + 跨租户拒绝 + 过期清理。"""
from __future__ import annotations

import asyncio
import io
import time

from sqlalchemy import select

from app.agents.tools.base import ToolContext
from app.agents.tools.file_tools import ConvertFileToTool, GenerateFileTool, ReadFileTool
from app.core.db import AsyncSessionLocal, init_models
from app.ingest.storage import get_storage
from app.models import Artifact, Tenant, User
from app.services.permission import PrincipalSet


async def _setup():
    await init_models()
    async with AsyncSessionLocal() as db:
        t = (await db.execute(select(Tenant).where(Tenant.slug == "ft"))).scalar_one_or_none()
        if not t:
            t = Tenant(name="FT", slug="ft"); db.add(t); await db.flush()
        u = (await db.execute(select(User).where(User.username == "ft_u"))).scalar_one_or_none()
        if not u:
            u = User(tenant_id=t.id, username="ft_u", password_hash="x"); db.add(u); await db.flush()
        await db.commit()
        return {"tid": t.id, "uid": u.id}


async def _run():
    d = await _setup()
    ps = PrincipalSet(user_id=d["uid"], tenant_id=d["tid"])

    # 1) 生成 docx
    async with AsyncSessionLocal() as db:
        ctx = ToolContext(db=db, ps=ps, tenant_id=d["tid"], user_id=d["uid"], conversation_id=1)
        res = await GenerateFileTool().execute(
            {"filename": "报告.docx", "format": "docx", "content": "# 标题\n正文段落"}, ctx)
        assert not res.is_error, res.content
        aid = res.data["artifact_id"]
        fk = res.data["file_key"]
        assert fk.startswith(f"{d['tid']}/"), fk
        # 文件落盘
        assert get_storage().path(fk).is_file()
        # Artifact 落库
        art = await db.get(Artifact, aid)
        assert art and art.file_name == "报告.docx" and art.source == "generated"

    # 2) 生成 xlsx（rows）
    async with AsyncSessionLocal() as db:
        ctx = ToolContext(db=db, ps=ps, tenant_id=d["tid"], user_id=d["uid"])
        res = await GenerateFileTool().execute(
            {"filename": "数据.xlsx", "format": "xlsx", "rows": [["a", "b"], [1, 2]]}, ctx)
        assert not res.is_error
        assert get_storage().path(res.data["file_key"]).is_file()

    # 3) 生成 pdf
    async with AsyncSessionLocal() as db:
        ctx = ToolContext(db=db, ps=ps, tenant_id=d["tid"], user_id=d["uid"])
        res = await GenerateFileTool().execute(
            {"filename": "文档.pdf", "format": "pdf", "content": "中文内容测试"}, ctx)
        assert not res.is_error, res.content

    # 4) 跨租户读取被拒
    async with AsyncSessionLocal() as db:
        ctx = ToolContext(db=db, ps=ps, tenant_id=d["tid"], user_id=d["uid"])
        res = await ReadFileTool().run({"file_key": "99999/abc.txt"}, ctx)
        assert res.is_error and "无权" in res.content, res.content

    # 5) xlsx→csv 转换
    async with AsyncSessionLocal() as db:
        ctx = ToolContext(db=db, ps=ps, tenant_id=d["tid"], user_id=d["uid"])
        gen = await GenerateFileTool().execute(
            {"filename": "src.xlsx", "format": "xlsx", "rows": [["x", "y"], [3, 4]]}, ctx)
        src_key = gen.data["file_key"]
        conv = await ConvertFileToTool().execute({"file_key": src_key, "target_ext": "csv"}, ctx)
        assert not conv.is_error, conv.content
        content = get_storage().read(conv.data["file_key"]).decode("utf-8")
        assert "x" in content and "3" in content, content

    # 6) 改写 docx：查找替换 + 追加
    from app.agents.tools.file_tools import EditFileTool

    async with AsyncSessionLocal() as db:
        ctx = ToolContext(db=db, ps=ps, tenant_id=d["tid"], user_id=d["uid"])
        gen = await GenerateFileTool().execute(
            {"filename": "草稿.docx", "format": "docx", "content": "# 草稿\n旧标题在这里\n"}, ctx)
        ed = await EditFileTool().execute(
            {"file_key": gen.data["file_key"], "replace": [{"find": "旧标题", "replace": "新标题"}],
             "instructions": "补充说明"}, ctx)
        assert not ed.is_error, ed.content
        import docx as _docx

        edoc = _docx.Document(io.BytesIO(get_storage().read(ed.data["file_key"])))
        texts = "\n".join(p.text for p in edoc.paragraphs)
        assert "新标题" in texts and "旧标题" not in texts, texts
        assert "补充说明" in texts

    # 7) 改写 xlsx：单元格写入 + 追加行
    async with AsyncSessionLocal() as db:
        ctx = ToolContext(db=db, ps=ps, tenant_id=d["tid"], user_id=d["uid"])
        gen = await GenerateFileTool().execute(
            {"filename": "表.xlsx", "format": "xlsx", "rows": [["名称", "值"], ["甲", 1]]}, ctx)
        ed = await EditFileTool().execute(
            {"file_key": gen.data["file_key"], "cells": {"B2": 99}, "append_rows": [["乙", 2]]}, ctx)
        assert not ed.is_error, ed.content
        import openpyxl as _x

        ws = _x.load_workbook(io.BytesIO(get_storage().read(ed.data["file_key"]))).active
        assert ws["B2"].value == 99 and ws["A3"].value == "乙"

    # 8) 过期清理
    async with AsyncSessionLocal() as db:
        db.add(Artifact(tenant_id=d["tid"], user_id=d["uid"], file_key=f"{d['tid']}/expired.bin",
                        file_name="expired.bin", size=1, source="generated",
                        expires_at=int(time.time() * 1000) - 1000))
        await db.commit()
    from app.tasks.artifact_tasks import cleanup_expired_artifacts

    async with AsyncSessionLocal() as db:
        n = await cleanup_expired_artifacts(db)
        assert n >= 1, n
    print("OK file_tools")


def test_file_tools():
    loop = asyncio.new_event_loop()
    try:
        asyncio.set_event_loop(loop)
        loop.run_until_complete(_run())
    finally:
        loop.close()
        asyncio.set_event_loop(None)


async def _run_office():
    """办公增强渲染：docx 封面/页眉/页码/表格样式；xlsx 多表/格式/图表/冻结。"""
    from app.agents.tools.office_render import render_docx, render_xlsx

    # docx：封面 + 页眉 + 页码域 + 表头底纹 + 中日韩字体
    md = "# 季度报告\n\n## 指标\n\n| 指标 | 值 |\n| --- | --- |\n| 营收 | 100 |\n"
    data = render_docx(md, meta={"title": "季度报告", "author": "张三", "cover": True,
                                 "header": "内部资料", "toc": True})
    assert data[:2] == b"PK", "docx 应为 zip 容器"
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(io.BytesIO(data))
    assert d.tables and d.tables[0].rows[0].cells[0].text == "指标"
    shd = d.tables[0].rows[0].cells[0]._tc.tcPr.find(qn("w:shd"))
    assert shd is not None and shd.get(qn("w:fill")) == "1E3A8A", "表头应有底纹"
    assert "PAGE" in d.sections[0].footer.paragraphs[0]._p.xml, "页脚应含 PAGE 域"
    assert d.sections[0].header.paragraphs[0].text == "内部资料"
    assert any("TOC" in p._p.xml for p in d.paragraphs), "应含目录域"
    rfonts = d.styles["Normal"].element.get_or_add_rPr().find(qn("w:rFonts"))
    assert rfonts is not None and rfonts.get(qn("w:eastAsia")), "应设置中文字体"

    # xlsx：多工作表 + 数字格式 + 冻结 + 图表（显式 data）
    sheets = [
        {"name": "明细", "rows": [["区域", "金额"], ["华东", 101570.0], ["华南", 34314.0]],
         "number_formats": {"B": "#,##0.00"},
         "charts": [{"type": "bar", "title": "各区域金额", "anchor": "D2",
                     "data": {"categories": ["华东", "华南"], "series": [{"name": "金额", "values": [101570, 34314]}]}}]},
        {"name": "汇总", "rows": [["部门", "达成率"], ["销售", 0.92]], "number_formats": {"B": "0.0%"}},
    ]
    xdata = render_xlsx("", sheets=sheets)
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xdata))
    assert wb.sheetnames == ["明细", "汇总"], wb.sheetnames
    ws = wb["明细"]
    assert ws.freeze_panes == "A2", "应冻结首行"
    assert ws["A1"].font.bold and ws["A1"].fill.fgColor.rgb.endswith("1E3A8A")
    assert ws["B2"].number_format == "#,##0.00"
    assert len(ws._charts) == 1, "应写入 1 个图表"
    assert wb["汇总"]["B2"].number_format == "0.0%"

    # xlsx：图表引用现有表格（不传 data）
    xdata2 = render_xlsx("", rows=[["区域", "金额"], ["华东", 1], ["华南", 2]],
                         charts=[{"type": "line", "title": "趋势", "anchor": "D2"}])
    ws2 = openpyxl.load_workbook(io.BytesIO(xdata2)).active
    assert len(ws2._charts) == 1, "引用现有表格也应生成图表"

    # pdf：封面 + 页眉 + 页码
    from app.agents.tools.office_render import render_pdf

    pdata = render_pdf("# 报告\n\n| 列 | 值 |\n| --- | --- |\n| A | 1 |\n",
                       meta={"title": "报告", "author": "李四", "cover": True, "header": "内部资料"})
    assert pdata[:5] == b"%PDF-", "应为 PDF"
    import pymupdf

    pdf = pymupdf.open(stream=pdata, filetype="pdf")
    assert pdf.page_count >= 2, "封面 + 正文应 ≥2 页"
    text = "\n".join(p.get_text() for p in pdf)
    from app.agents.tools.office_render import _pdf_register_fonts

    reg, _ = _pdf_register_fonts()
    hint = (
        f"\n>>> PDF 中文字体未注册（当前使用 {reg}），中文被渲染为占位符。"
        "\n>>> 本机：装中文字体；CI/Docker：apt-get install -y fonts-wqy-zenhei"
        "\n>>> 注意必须 TrueType 轮廓（fonts-noto-cjk 是 CFF，reportlab 会拒绝）。"
    )
    assert reg != "Helvetica", "未找到可用中文字体，PDF 中文会乱码。" + hint
    assert "报告" in text and "第 2 页" in text, text[:200] + hint


def test_office_render():
    asyncio.new_event_loop().run_until_complete(_run_office())


def test_content_disposition_handles_non_ascii():
    """文件名含中文时 Content-Disposition 必须可 latin-1 编码（否则响应 500）。"""
    from app.core.http_utils import content_disposition

    for name in ("pelican-bicycle.html", "鹈鹕骑自行车.html", "合并导出-3篇.pdf", 'a b"c.txt'):
        val = content_disposition(name, inline=True)
        val.encode("latin-1")  # 复现 starlette 头部编码，不能再抛 UnicodeEncodeError
        assert val.startswith("inline;")
        assert "filename*=UTF-8''" in val
    assert content_disposition("中文.docx").startswith("attachment;")
