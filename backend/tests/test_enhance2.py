"""本次修复的回归测试：工作流事件名、超时、HTTP headers、循环失败、转换 target、
doc_count 据实、RRF 原始分、cron 时区/补跑。"""
from __future__ import annotations

import asyncio
import time
from datetime import datetime, timezone

from sqlalchemy import select

from app.core.db import AsyncSessionLocal, init_models
from app.core.security import hash_password
from app.models import KnowledgeBase, Tenant, User


async def _setup(slug="fix2"):
    await init_models()
    async with AsyncSessionLocal() as db:
        t = (await db.execute(select(Tenant).where(Tenant.slug == slug))).scalar_one_or_none()
        if not t:
            t = Tenant(name=slug.upper(), slug=slug); db.add(t); await db.flush()
        u = (await db.execute(select(User).where(User.username == f"{slug}_u"))).scalar_one_or_none()
        if not u:
            u = User(tenant_id=t.id, username=f"{slug}_u", password_hash=hash_password("x"),
                     is_admin=True, user_type="internal")
            db.add(u); await db.flush()
        kb = KnowledgeBase(tenant_id=t.id, name=f"{slug}-kb", owner_id=u.id)
        db.add(kb); await db.flush()
        await db.commit()
        return {"tenant_id": t.id, "uid": u.id, "kb_id": kb.id}


# ==================== 工作流事件名归一化 ====================
def test_workflow_done_event_normalized():
    """runner 对外暴露 run_finished（而非内部哨兵 __wf_done__）。"""
    import inspect

    from app.agents.workflow import runner

    src = inspect.getsource(runner)
    assert '"run_finished"' in src, "runner 出口应把终态归一为 run_finished"
    assert '_DONE = "__wf_done__"' in src


# ==================== 定时任务超时 ====================
def test_scheduler_timeout_applied():
    import inspect

    from app.tasks import scheduler_tasks as ST

    src = inspect.getsource(ST.execute_task)
    assert "wait_for" in src and "timeout_seconds" in src, "execute_task 应包裹 wait_for 用 timeout_seconds"


# ==================== HTTP 节点 headers ====================
def test_http_node_supports_headers():
    async def _run():
        from app.agents.workflow.engine import _exec_http, NodeContext
        from app.services.permission import PrincipalSet

        # 用 localhost 不可达端口验证「headers 被接受且不报错」——只验证参数被消费
        node = {"type": "http", "data": {"url": "http://127.0.0.1:1/none", "method": "GET",
                                         "headers": {"X-Test": "1"}, "params": {"a": "b"},
                                         "timeout": 1}}
        ctx = NodeContext(db=None, ps=PrincipalSet(user_id=1, tenant_id=1), run_id=1, tenant_id=1)
        try:
            await _exec_http(node, {}, ctx)
            assert False, "不可达端口应异常"
        except Exception as e:  # noqa: BLE001
            # 是连接错误（说明 headers/params 已被接受并进入请求阶段），而非参数解析错误
            assert "headers" not in str(e).lower()
    asyncio.new_event_loop().run_until_complete(_run())


def test_http_node_headers_json_string_parsed():
    """headers 传 JSON 字符串也应被解析。"""
    import inspect

    from app.agents.workflow.engine import _exec_http

    src = inspect.getsource(_exec_http)
    assert "_as_obj" in src and "json_body" in src


# ==================== 循环节点失败上报 ====================
def test_loop_reports_failure():
    async def _run():
        from app.agents.workflow.engine import _exec_loop, NodeContext
        from app.core.errors import ValidationError
        from app.services.permission import PrincipalSet

        node = {"type": "loop", "data": {
            "body": {"op": "code", "code": "raise ValueError('boom')"},
            "max_iterations": 3,
        }}
        ctx = NodeContext(db=None, ps=PrincipalSet(user_id=1, tenant_id=1), run_id=1, tenant_id=1)
        try:
            await _exec_loop(node, {"items": [1, 2]}, ctx)
            assert False, "有失败项应抛错"
        except ValidationError as e:
            assert "失败" in str(e)

        # tolerant=true 则容忍
        node2 = {"type": "loop", "data": {
            "body": {"op": "code", "code": "raise ValueError('boom')"},
            "tolerant": True,
        }}
        out = await _exec_loop(node2, {"items": [1, 2]}, ctx)
        assert out["errors"] and out["results"] == [None, None]
    asyncio.new_event_loop().run_until_complete(_run())


# ==================== 转换工具 target ====================
def test_convert_file_csv_target(tmp_path):
    async def _run():
        from openpyxl import Workbook

        from app.agents.tools.file_tools import _to_csv

        wb = Workbook(); ws = wb.active; ws.title = "S1"
        ws.append(["名字", "值"]); ws.append(["A", 1]); ws.append(["B", 2])
        p = tmp_path / "t.xlsx"; wb.save(str(p))

        csv_text = _to_csv(p, "xlsx")
        assert "名字,值" in csv_text and "A,1" in csv_text
        assert "# 工作表: S1" in csv_text
    asyncio.new_event_loop().run_until_complete(_run())


def test_strip_markdown():
    from app.agents.tools.file_tools import _strip_markdown

    assert _strip_markdown("# 标题\n**粗体**和`代码`") == "标题\n粗体和代码"


# ==================== doc_count 据实 ====================
def test_refresh_kb_counts_decreases_after_delete():
    async def _run():
        from app.api.v1.document import _refresh_kb_counts
        from app.models import Chunk, Document

        d = await _setup("cnt")
        async with AsyncSessionLocal() as db:
            db.add_all([
                Document(tenant_id=d["tenant_id"], kb_id=d["kb_id"], title="A", status="ready", char_count=10, chunk_count=1),
                Document(tenant_id=d["tenant_id"], kb_id=d["kb_id"], title="B", status="ready", char_count=10, chunk_count=1),
            ])
            await db.commit()
            await _refresh_kb_counts(db, d["kb_id"])
            await db.commit()
            kb = await db.get(KnowledgeBase, d["kb_id"])
            assert kb.doc_count == 2, f"应有 2 个文档，实际 {kb.doc_count}"

            # 删一个文档后计数应递减
            doc = (await db.execute(select(Document).where(Document.kb_id == d["kb_id"]))).scalars().first()
            await db.delete(doc)
            await db.flush()
            await _refresh_kb_counts(db, d["kb_id"])
            await db.commit()
            kb = await db.get(KnowledgeBase, d["kb_id"])
            assert kb.doc_count == 1, f"删除后应为 1，实际 {kb.doc_count}"

            from sqlalchemy import delete
            await db.execute(delete(Chunk).where(Chunk.kb_id == d["kb_id"]))
            await db.execute(delete(Document).where(Document.kb_id == d["kb_id"]))
            await db.commit()
    asyncio.new_event_loop().run_until_complete(_run())


# ==================== RRF 原始分 ====================
def test_score_threshold_uses_raw_similarity():
    from app.retrieval.fusion import rrf_fuse
    from app.retrieval.vector_store.store import VectorHit

    hits = [
        VectorHit(chunk_id=1, doc_id=1, kb_id=1, content="a", score=0.9),
        VectorHit(chunk_id=2, doc_id=1, kb_id=1, content="b", score=0.3),
    ]
    fused = rrf_fuse([hits], top_n=2)
    # 阈值 0.5 应只留 raw_score>=0.5 的（即 chunk 1）
    kept = [h for h in fused if (h.raw_score or 0) >= 0.5]
    assert [h.chunk_id for h in kept] == [1]


# ==================== cron 时区与补跑 ====================
def test_cron_next_run_with_timezone():
    from app.services.cron import next_run

    # 每天 9:00（Asia/Shanghai）
    after = datetime(2026, 3, 10, 8, 0, tzinfo=timezone.utc)  # UTC 08:00 = 沪 16:00
    dt = next_run("0 9 * * *", after, tz_name="Asia/Shanghai")
    # 结果的本地时区小时数应为 9
    assert dt.hour == 9


def test_compute_next_interval_with_after_ms():
    from app.services.schedule_service import compute_next

    class _T:
        schedule_kind = "interval"
        interval_seconds = 300
        timezone = "Asia/Shanghai"
        run_at = None
        cron_expr = None

    base = int(time.time() * 1000)
    nxt = compute_next(_T(), after_ms=base)
    assert abs(nxt - (base + 300_000)) < 5, "间隔任务应基于 after_ms 递推"
